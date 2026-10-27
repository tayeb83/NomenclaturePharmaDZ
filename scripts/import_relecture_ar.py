#!/usr/bin/env python3
"""
Réimporte le classeur de relecture (scripts/export_relecture_ar.py) dans
data/ansm/interactions_ar.json.

Pour chaque ligne de l'onglet « À relire » :
  - « Correction AR » remplie : la correction remplace la traduction, et le texte est marqué relu ;
  - Statut « Relu » : le texte est marqué relu, tel quel ;
  - Statut « À relire » : le texte n'est pas relu (s'il l'était, il ne l'est plus).
Les corrections sont contrôlées comme à la traduction : de l'arabe, les mêmes nombres, les mêmes paragraphes que le français.
Une ligne refusée est signalée et laissée telle quelle ; le reste est importé.

Les textes relus sont listés dans la clé « relus » du JSON. scripts/import_ansm_traductions_ar.py en déduit les interactions
à marquer « traduction relue » (risque et conduite relus).

Usage :
    python scripts/import_relecture_ar.py data/ansm/relecture_ar.xlsx            # simulation
    python scripts/import_relecture_ar.py data/ansm/relecture_ar.xlsx --apply

Prérequis : pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CATEGORIES = {"Risque": "risques", "Conduite": "conduites", "Partenaire": "partenaires"}
NOMBRE = re.compile(r"\d+(?:[.,]\d+)?")


def controle(fr: str, ar: str, categorie: str) -> list[str]:
    erreurs = []
    if not re.search(r"[؀-ۿ]", ar):
        erreurs.append("aucun caractère arabe")
    if categorie != "partenaires":
        manquants = [n for n in NOMBRE.findall(fr) if n not in NOMBRE.findall(ar)]
        if manquants:
            erreurs.append("nombres absents : " + ", ".join(manquants))
        if fr.count("\n") != ar.count("\n"):
            erreurs.append(f"paragraphes : {fr.count(chr(10)) + 1} en français, {ar.count(chr(10)) + 1} en arabe")
    return erreurs


def texte(v) -> str:
    return str(v).replace("\r\n", "\n").strip() if v is not None else ""


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("classeur", type=Path)
    ap.add_argument("--ar", type=Path, default=ROOT / "data/ansm/interactions_ar.json")
    ap.add_argument("--apply", action="store_true", help="écrire le JSON (sinon : simulation)")
    args = ap.parse_args()

    ar = json.loads(args.ar.read_text(encoding="utf8"))
    relus = {k: set(v) for k, v in ar.get("relus", {}).items()}
    for cat in CATEGORIES.values():
        relus.setdefault(cat, set())

    ws = load_workbook(args.classeur, data_only=True)["À relire"]
    entetes = {texte(c.value): i for i, c in enumerate(ws[1])}
    for nom in ("Type", "Texte français (ANSM)", "Traduction arabe", "Statut", "Correction AR"):
        if nom not in entetes:
            print(f"Colonne absente : {nom}")
            return 2

    corriges = marques = retires = inconnus = 0
    refus: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        libelle, fr = texte(row[entetes["Type"]]), texte(row[entetes["Texte français (ANSM)"]])
        if not fr:
            continue
        cat = CATEGORIES.get(libelle)
        # l'export écrit les paragraphes avec « \n » : on compare sans tenir compte des fins de ligne
        cible = next((k for k in (ar.get(cat, {}) if cat else {}) if k.replace("\r\n", "\n").strip() == fr), None)
        if cat is None or cible is None:
            inconnus += 1
            continue
        statut, correction = texte(row[entetes["Statut"]]), texte(row[entetes["Correction AR"]])

        if correction:
            erreurs = controle(cible, correction, cat)
            if erreurs:
                refus.append(f"{libelle} « {fr[:60]}… » : {' ; '.join(erreurs)}")
                continue
            if correction != ar[cat][cible]:
                ar[cat][cible] = correction
                corriges += 1
            if cible not in relus[cat]:
                relus[cat].add(cible)
                marques += 1
        elif statut == "Relu":
            if cible not in relus[cat]:
                relus[cat].add(cible)
                marques += 1
        elif statut == "À relire" and cible in relus[cat]:
            relus[cat].discard(cible)
            retires += 1

    ar["relus"] = {cat: sorted(v) for cat, v in relus.items()}
    print(f"{corriges} traductions corrigées, {marques} textes marqués relus, {retires} retirés des relus, {inconnus} lignes non reconnues")
    print("relus : " + ", ".join(f"{len(ar['relus'][c])} {c}" for c in CATEGORIES.values()))
    for r in refus:
        print("  refusé :", r)
    if args.apply:
        args.ar.write_text(json.dumps(ar, ensure_ascii=False, indent=1) + "\n", encoding="utf8")
        print(f"Écrit : {args.ar}")
    else:
        print("Simulation : rien n'est écrit. Relancez avec --apply.")
    return 1 if refus else 0


if __name__ == "__main__":
    raise SystemExit(main())
